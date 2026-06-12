"""
anonymize_file — Conversion et anonymisation de fichiers (TXT, DOCX, PDF).

Ce module convertit les fichiers DOCX et PDF en texte brut (UTF-8),
puis anonymise le texte via le moteur NLP partagé de ``api.nlp_engine``.
Le résultat est toujours un fichier ``.txt`` anonymisé.

Formats supportés en entrée
---------------------------
- **.txt** : anonymisation directe
- **.docx** : extraction du texte → anonymisation
- **.pdf** : extraction du texte → anonymisation
"""

from __future__ import annotations

import io

import fitz  # PyMuPDF
from docx import Document

from api.nlp_engine import anonymize_text


# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

MAX_FILE_SIZE: int = int(10 * 1024 * 1024)
"""Taille maximale d'un fichier uploadé (10 Mo)."""


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------

def _process_txt(content: bytes) -> bytes:
    """Anonymise un fichier texte brut.

    Parameters
    ----------
    content : bytes
        Contenu brut du fichier ``.txt``.

    Returns
    -------
    bytes
        Contenu anonymisé encodé en UTF-8.
    """
    text = content.decode("utf-8", errors="replace")
    return anonymize_text(text).encode("utf-8")


# ---------------------------------------------------------------------------
# DOCX → TXT
# ---------------------------------------------------------------------------

def _convert_docx_to_txt(content: bytes) -> bytes:
    """Extrait le texte brut d'un fichier DOCX.

    Parcourt les paragraphes du corps principal ainsi que les cellules
    de tableaux, et retourne le tout en texte UTF-8.

    Parameters
    ----------
    content : bytes
        Contenu brut du fichier ``.docx``.

    Returns
    -------
    bytes
        Texte brut encodé en UTF-8.
    """
    doc = Document(io.BytesIO(content))
    lines: list[str] = []

    # Paragraphes principaux
    for paragraph in doc.paragraphs:
        lines.append(paragraph.text)

    # Cellules de tableaux
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)

    return "\n".join(lines).encode("utf-8")


def _process_docx(content: bytes) -> bytes:
    """Convertit un DOCX en texte brut, puis anonymise.

    Parameters
    ----------
    content : bytes
        Contenu brut du fichier ``.docx``.

    Returns
    -------
    bytes
        Texte anonymisé encodé en UTF-8.
    """
    txt_bytes = _convert_docx_to_txt(content)
    return _process_txt(txt_bytes)


# ---------------------------------------------------------------------------
# PDF → TXT
# ---------------------------------------------------------------------------

def _convert_pdf_to_txt(content: bytes) -> bytes:
    """Extrait le texte brut d'un fichier PDF.

    Parcourt chaque page et concatène le texte extrait,
    séparé par des sauts de ligne.

    Parameters
    ----------
    content : bytes
        Contenu brut du fichier ``.pdf``.

    Returns
    -------
    bytes
        Texte brut encodé en UTF-8.
    """
    doc = fitz.open(stream=content, filetype="pdf")
    pages: list[str] = []

    for page in doc:
        text = page.get_text("text")
        if text.strip():
            pages.append(text)

    doc.close()
    return "\n".join(pages).encode("utf-8")


def _process_pdf(content: bytes) -> bytes:
    """Convertit un PDF en texte brut, puis anonymise.

    Parameters
    ----------
    content : bytes
        Contenu brut du fichier ``.pdf``.

    Returns
    -------
    bytes
        Texte anonymisé encodé en UTF-8.
    """
    txt_bytes = _convert_pdf_to_txt(content)
    return _process_txt(txt_bytes)
