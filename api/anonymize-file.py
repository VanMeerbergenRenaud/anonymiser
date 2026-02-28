from http.server import BaseHTTPRequestHandler
import json
import io
import re
import tempfile
import os
import cgi

import fitz  # PyMuPDF
from docx import Document

# Reuse the shared NLP engine from anonymize-text
import spacy
from presidio_analyzer import AnalyzerEngine, PatternRecognizer, Pattern, RecognizerRegistry
from presidio_analyzer.nlp_engine import NlpEngineProvider
from presidio_anonymizer import AnonymizerEngine
from presidio_anonymizer.entities import OperatorConfig

# ---------------------------------------------------------------------------
# NLP + Presidio initialisation (same as anonymize-text, module-level cache)
# ---------------------------------------------------------------------------

def _build_analyzer() -> AnalyzerEngine:
    configuration = {
        "nlp_engine_name": "spacy",
        "models": [{"lang_code": "fr", "model_name": "fr_core_news_sm"}],
    }
    provider = NlpEngineProvider(nlp_configuration=configuration)
    nlp_engine = provider.create_engine()

    nir_pattern = Pattern(
        name="nir_pattern",
        regex=r"\b[12]\s?\d{2}\s?\d{2}\s?\d{2}\s?\d{3}\s?\d{3}\s?\d{2}\b",
        score=0.9,
    )
    nir_recognizer = PatternRecognizer(
        supported_entity="FR_NIR",
        name="French NIR Recognizer",
        patterns=[nir_pattern],
        supported_language="fr",
    )

    iban_pattern = Pattern(
        name="iban_fr_pattern",
        regex=r"\b[A-Z]{2}\d{2}\s?\d{4}\s?\d{4}\s?\d{4}\s?\d{4}\s?\d{4}\s?\d{0,3}\b",
        score=0.85,
    )
    iban_recognizer = PatternRecognizer(
        supported_entity="IBAN_CODE",
        name="IBAN Recognizer",
        patterns=[iban_pattern],
        supported_language="fr",
    )

    phone_pattern = Pattern(
        name="fr_phone_pattern",
        regex=r"\b(?:(?:\+33|0033)\s?[1-9](?:[\s.-]?\d{2}){4}|0[1-9](?:[\s.-]?\d{2}){4})\b",
        score=0.85,
    )
    phone_recognizer = PatternRecognizer(
        supported_entity="PHONE_NUMBER",
        name="French Phone Recognizer",
        patterns=[phone_pattern],
        supported_language="fr",
    )

    postal_pattern = Pattern(
        name="fr_postal_pattern",
        regex=r"\b\d{5}\b",
        score=0.4,
    )
    postal_recognizer = PatternRecognizer(
        supported_entity="FR_POSTAL_CODE",
        name="French Postal Code Recognizer",
        patterns=[postal_pattern],
        supported_language="fr",
        context=["code postal", "cp", "adresse", "domicilié", "résidant", "habite", "situé"],
    )

    email_pattern = Pattern(
        name="email_pattern",
        regex=r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b",
        score=0.9,
    )
    email_recognizer = PatternRecognizer(
        supported_entity="EMAIL_ADDRESS",
        name="Email Recognizer FR",
        patterns=[email_pattern],
        supported_language="fr",
    )

    registry = RecognizerRegistry()
    registry.load_predefined_recognizers(nlp_engine=nlp_engine, languages=["fr"])
    registry.add_recognizer(nir_recognizer)
    registry.add_recognizer(iban_recognizer)
    registry.add_recognizer(phone_recognizer)
    registry.add_recognizer(postal_recognizer)
    registry.add_recognizer(email_recognizer)

    return AnalyzerEngine(
        nlp_engine=nlp_engine,
        registry=registry,
        supported_languages=["fr"],
    )


analyzer = _build_analyzer()
anonymizer_engine = AnonymizerEngine()

ENTITY_LABELS: dict[str, str] = {
    "PERSON": "PERSONNE",
    "LOCATION": "LIEU",
    "DATE_TIME": "DATE",
    "EMAIL_ADDRESS": "EMAIL",
    "PHONE_NUMBER": "TÉLÉPHONE",
    "IBAN_CODE": "IBAN",
    "FR_NIR": "NIR",
    "FR_POSTAL_CODE": "CODE_POSTAL",
    "NRP": "NATIONALITÉ",
    "CREDIT_CARD": "CARTE_BANCAIRE",
    "IP_ADDRESS": "ADRESSE_IP",
    "URL": "URL",
}

MAX_FILE_SIZE = 4.5 * 1024 * 1024  # 4.5 MB


def _get_label(entity_type: str) -> str:
    return ENTITY_LABELS.get(entity_type, entity_type)


# ---------------------------------------------------------------------------
# TXT processing
# ---------------------------------------------------------------------------

def _process_txt(content: bytes) -> bytes:
    text = content.decode("utf-8", errors="replace")
    results = analyzer.analyze(text=text, language="fr")
    operators = {}
    for entity_type in set(r.entity_type for r in results):
        operators[entity_type] = OperatorConfig(
            "replace", {"new_value": f"[{_get_label(entity_type)}]"}
        )
    anonymized = anonymizer_engine.anonymize(text=text, analyzer_results=results, operators=operators)
    return anonymized.text.encode("utf-8")


# ---------------------------------------------------------------------------
# DOCX processing — replace sensitive text while preserving paragraphs
# ---------------------------------------------------------------------------

def _process_docx(content: bytes) -> bytes:
    doc = Document(io.BytesIO(content))

    for paragraph in doc.paragraphs:
        full_text = paragraph.text
        if not full_text.strip():
            continue

        results = analyzer.analyze(text=full_text, language="fr")
        if not results:
            continue

        # Sort results from end to start to preserve indices
        results = sorted(results, key=lambda r: r.start, reverse=True)

        # Work on the full text, then redistribute to runs
        new_text = full_text
        for result in results:
            label = _get_label(result.entity_type)
            new_text = new_text[: result.start] + f"[{label}]" + new_text[result.end:]

        # Replace runs: clear all runs and set text in the first run
        if paragraph.runs:
            # Preserve formatting of the first run
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""

    # Also process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    full_text = paragraph.text
                    if not full_text.strip():
                        continue
                    results = analyzer.analyze(text=full_text, language="fr")
                    if not results:
                        continue
                    results = sorted(results, key=lambda r: r.start, reverse=True)
                    new_text = full_text
                    for result in results:
                        label = _get_label(result.entity_type)
                        new_text = new_text[: result.start] + f"[{label}]" + new_text[result.end:]
                    if paragraph.runs:
                        paragraph.runs[0].text = new_text
                        for run in paragraph.runs[1:]:
                            run.text = ""

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF processing — redact with black rectangles over detected entities
# ---------------------------------------------------------------------------

def _process_pdf(content: bytes) -> bytes:
    doc = fitz.open(stream=content, filetype="pdf")

    for page in doc:
        text_page = page.get_text("text")
        if not text_page.strip():
            continue

        results = analyzer.analyze(text=text_page, language="fr")
        if not results:
            continue

        for result in results:
            sensitive_text = text_page[result.start: result.end]
            # Search for all instances of this text on the page
            text_instances = page.search_for(sensitive_text)
            for inst in text_instances:
                # Add a black redaction annotation
                annot = page.add_redact_annot(inst, fill=(0, 0, 0))
        
        # Apply all redactions on this page
        page.apply_redactions()

    buf = io.BytesIO()
    doc.save(buf, garbage=4, deflate=True)
    doc.close()
    return buf.getvalue()


# ---------------------------------------------------------------------------
# HTTP handler
# ---------------------------------------------------------------------------

def _parse_multipart(handler):
    """Parse multipart form data and return the file info."""
    content_type = handler.headers.get("Content-Type", "")
    
    if "multipart/form-data" not in content_type:
        return None, None, None
    
    # Extract boundary from content-type
    boundary = None
    for part in content_type.split(";"):
        part = part.strip()
        if part.startswith("boundary="):
            boundary = part[len("boundary="):]
            break
    
    if not boundary:
        return None, None, None
    
    content_length = int(handler.headers.get("Content-Length", 0))
    body = handler.rfile.read(content_length)
    
    # Parse the multipart body
    boundary_bytes = boundary.encode()
    parts = body.split(b"--" + boundary_bytes)
    
    for part in parts:
        if b"Content-Disposition" not in part:
            continue
        
        # Split headers from body
        header_end = part.find(b"\r\n\r\n")
        if header_end == -1:
            continue
        
        headers_raw = part[:header_end].decode("utf-8", errors="replace")
        file_data = part[header_end + 4:]
        
        # Remove trailing \r\n
        if file_data.endswith(b"\r\n"):
            file_data = file_data[:-2]
        
        # Extract filename
        filename = None
        for line in headers_raw.split("\r\n"):
            if "filename=" in line:
                match = re.search(r'filename="([^"]*)"', line)
                if match:
                    filename = match.group(1)
                    break
        
        if filename and file_data:
            # Detect content type from filename
            ext = os.path.splitext(filename)[1].lower()
            return filename, ext, file_data
    
    return None, None, None


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            content_length = int(self.headers.get("Content-Length", 0))
            if content_length > MAX_FILE_SIZE:
                self._error(413, "Le fichier est trop volumineux (max 4.5 Mo).")
                return

            filename, ext, file_data = _parse_multipart(self)

            if not filename or not file_data:
                self._error(400, "Aucun fichier valide reçu.")
                return

            # Process based on file extension
            if ext == ".txt":
                result_data = _process_txt(file_data)
                content_type = "text/plain; charset=utf-8"
            elif ext == ".docx":
                result_data = _process_docx(file_data)
                content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif ext == ".pdf":
                result_data = _process_pdf(file_data)
                content_type = "application/pdf"
            else:
                self._error(
                    400,
                    f"Format non supporté : {ext}. Formats acceptés : .txt, .docx, .pdf",
                )
                return

            # Build anonymized filename
            name_without_ext = os.path.splitext(filename)[0]
            anon_filename = f"{name_without_ext}_anonymise{ext}"

            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header(
                "Content-Disposition",
                f'attachment; filename="{anon_filename}"',
            )
            self.send_header("Content-Length", str(len(result_data)))
            self.send_header("X-Filename", anon_filename)
            self.end_headers()
            self.wfile.write(result_data)

        except Exception as e:
            self._error(500, f"Erreur lors du traitement du fichier : {str(e)}")

    def _error(self, code: int, message: str):
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(json.dumps({"error": message}, ensure_ascii=False).encode("utf-8"))
