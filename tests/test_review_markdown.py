"""Mots à relire (``api.review``) et sortie Markdown des fichiers."""

import io
import re

from docx import Document

from api.anonymize_file import process_file, to_markdown
from api.nlp_engine import anonymize_with_review


def terms(text: str) -> dict[str, str]:
    return {t["term"]: t["reason"] for t in anonymize_with_review(text)[1]}


# ---------------------------------------------------------------------------
# Mots à relire
# ---------------------------------------------------------------------------

def test_review_lists_what_was_kept_on_purpose():
    found = terms(
        "Voir Cour eur. D.H., arrêt Michaud c. France du 6 décembre 2012. La SRL Bati-Meuse "
        "a livré les matériaux sur le chantier."
    )
    assert "Michaud" in found and "jurisprudence" in found["Michaud"]
    assert any("Bati-Meuse" in t for t in found)


def test_review_ignores_safe_and_masked_terms():
    found = terms(
        "Madame Sophie LEMAIRE, domiciliée rue des Écoles 14 à 4800 Verviers, a saisi le "
        "Tribunal de première instance de Liège, division Verviers, et la Cour de cassation. "
        "Vu le Code civil et la directive (UE) 2018/822."
    )
    assert not any(w in t for t in found for w in ("LEMAIRE", "Sophie", "Liège", "Code", "Cour"))
    assert not any("Écoles" in t or "Verviers" in t for t in found)


def test_review_flags_unknown_capitalized_word():
    text = "Les fonds ont été versés par Zrbqwtk au profit de la défenderesse."
    out, review = anonymize_with_review(text)
    if "Zrbqwtk" in out:  # non masqué : doit être signalé
        assert any(t["term"] == "Zrbqwtk" for t in review)


def test_review_ignores_sentence_starts_and_acronyms():
    found = terms("Partant, le recours est fondé. Considérant que le CIR et la TVA s'appliquent.")
    assert found == {}


def test_text_endpoint_returns_review():
    from api.index import app

    response = app.test_client().post("/api/anonymize_text", json={
        "text": "Voir l'arrêt Michaud c. France (Cour eur. D.H., 6 décembre 2012)."})
    data = response.get_json()
    assert isinstance(data["review"], list) and "anonymized" in data


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def _docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_word_headings_and_tables_become_markdown():
    def build(d):
        d.add_heading("Conclusions de synthèse", level=1)
        d.add_heading("Les faits", level=2)
        d.add_paragraph("Le contrat a été signé le 3 mars 2020.")
        d.add_paragraph("Il a été résilié le 4 avril 2021.")
        table = d.add_table(rows=3, cols=3)
        for r, row in enumerate([["Pièce", "Date", "Objet"], ["1", "", "Contrat"],
                                 ["2", "4 avril 2021", "Résiliation"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value

    out = process_file("c.docx", _docx(build)).content.decode("utf-8")
    assert "# Conclusions de synthèse\n" in out
    assert "\n## Les faits\n" in out
    assert "signé le 3 mars 2020.\n\nIl a été résilié" in out
    assert "| Pièce | Date | Objet |\n|---|---|---|\n| 1 |  | Contrat |\n| 2 | 4 avril 2021 | Résiliation |" in out


def test_single_column_table_is_plain_text():
    def build(d):
        table = d.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Cabinet d'avocats"
        table.cell(1, 0).text = "Avenue Louise 500"

    out = process_file("t.docx", _docx(build)).content.decode("utf-8")
    assert "|" not in out and "Cabinet d'avocats" in out


def test_to_markdown_changes_whitespace_only():
    text = "Titre\n\tLe tribunal statue.\n| a | b |\n|---|---|\n| c | d |\n\n\nFin."
    md = to_markdown(text)
    assert re.findall(r"\S+", md) == re.findall(r"\S+", text)
    assert "\n\nLe tribunal statue.\n\n| a | b |\n|---|---|\n| c | d |\n\nFin.\n" in md
