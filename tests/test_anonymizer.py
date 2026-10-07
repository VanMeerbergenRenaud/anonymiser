"""Tests de bout en bout : texte, fichiers (TXT, DOCX, PDF, images) et OCR.

Charge le modèle NER (CamemBERT ou spaCy) : plus lent que les tests de règles.
Les tests OCR sont ignorés si Tesseract n'est pas installé.
"""

import io

import fitz
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

from api import ocr
from api.anonymize_file import FileProcessingError, process_file
from api.nlp_engine import anonymize_text

needs_ocr = pytest.mark.skipif(not ocr.is_available(), reason="Tesseract non installé")

BELGIAN_TEXT = (
    "TRIBUNAL DE PREMIÈRE INSTANCE DE LIÈGE\n"
    "Monsieur Olivier DEPREZ, né à Seraing le 12 octobre 1978 (NN 78.10.12-345.67), "
    "domicilié Rue de la Province 45 boîte 2, 4100 Seraing, GSM 0472 11 22 33, "
    "olivier.deprez@exemple.be, compte BE68 5390 0754 7034.\n"
    "PAR CES MOTIFS, vu l'article 1184 de l'ancien Code civil, le tribunal "
    "condamne la société à payer à Monsieur DEPREZ la somme de 25.000 €."
)

SECRETS = [
    "Olivier", "DEPREZ", "12 octobre 1978", "78.10.12-345.67", "Province",
    "4100", "0472 11 22 33", "olivier.deprez", "BE68",
]


def _font(size: int):
    try:
        return ImageFont.truetype("DejaVuSans.ttf", size)
    except OSError:
        return ImageFont.load_default(size=size)


def text_png(lines: list[str], size=(1100, 300)) -> bytes:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((30, 30 + i * 50), line, fill="black", font=_font(32))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def test_belgian_text_is_anonymized():
    result = anonymize_text(BELGIAN_TEXT)
    for secret in SECRETS:
        assert secret not in result, secret
    # Intitulés et références juridiques conservés.
    assert "TRIBUNAL DE PREMIÈRE INSTANCE" in result
    assert "PAR CES MOTIFS" in result
    assert "Code civil" in result
    assert "[REGISTRE_NATIONAL]" in result
    assert "[ADRESSE]" in result


def test_same_person_gets_same_pseudonym():
    result = anonymize_text(
        "Monsieur Jean DUPONT a signé. Madame Sophie MARTIN a refusé. "
        "Plus tard, Monsieur DUPONT a relancé Madame MARTIN."
    )
    assert result.count("[PERSONNE_1]") == 2
    assert result.count("[PERSONNE_2]") == 2


def test_ocr_markers_are_preserved():
    text = f"{ocr.OCR_START}\nMadame Claire DUBOIS\n{ocr.OCR_END}"
    result = anonymize_text(text)
    assert result.startswith(ocr.OCR_START)
    assert result.endswith(ocr.OCR_END)
    assert "DUBOIS" not in result


def test_txt_windows_1252():
    content = "Madame Hélène GOFFIN, née à Huy le 2 août 1990.".encode("cp1252")
    result = process_file("note.txt", content).content.decode("utf-8")
    assert "GOFFIN" not in result and "2 août 1990" not in result
    assert "née à" in result


def test_unsupported_and_corrupted_files():
    with pytest.raises(FileProcessingError):
        process_file("ancien.doc", b"\xd0\xcf\x11\xe0")
    with pytest.raises(FileProcessingError):
        process_file("casse.docx", b"pas un zip")
    with pytest.raises(FileProcessingError):
        process_file("casse.pdf", b"pas un pdf")


def test_password_protected_pdf():
    pdf = fitz.open()
    pdf.new_page().insert_text((50, 50), "Secret")
    data = pdf.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    with pytest.raises(FileProcessingError):
        process_file("protege.pdf", data)


@needs_ocr
def test_docx_image_text_is_extracted_and_anonymized():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Étude Maître Isabelle HENRARD"
    doc.add_paragraph("Voir la capture :")
    doc.add_picture(io.BytesIO(text_png([
        "Titulaire : Nathalie VANDENBERGHE",
        "Compte : BE71 0961 2345 6769",
    ])), width=Inches(5))
    buffer = io.BytesIO()
    doc.save(buffer)

    result = process_file("bail.docx", buffer.getvalue())
    text = result.content.decode("utf-8")
    assert result.ocr_images == 1
    assert ocr.OCR_START in text and ocr.OCR_END in text
    assert "VANDENBERGHE" not in text and "BE71" not in text and "HENRARD" not in text


@needs_ocr
def test_scanned_pdf_is_read_by_ocr():
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(30, 40, 565, 300), stream=text_png([
        "Madame Claire DUBOIS",
        "Rue Saint-Gilles 120, 4000 Liège",
        "Registre national : 85.07.30-033.28",
    ], size=(1600, 400)))
    result = process_file("scan.pdf", pdf.tobytes())
    text = result.content.decode("utf-8")
    assert result.ocr_images == 1
    assert text.startswith(ocr.OCR_START)
    for secret in ("DUBOIS", "Saint-Gilles", "85.07.30-033.28"):
        assert secret not in text


@needs_ocr
def test_rotated_image_upload():
    img = Image.open(io.BytesIO(text_png([
        "Monsieur Pierre DELVAUX",
        "Avenue Louise 54 bte 3, 1050 Ixelles",
    ], size=(900, 200)))).rotate(90, expand=True)
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    text = process_file("attestation.png", buffer.getvalue()).content.decode("utf-8")
    assert ocr.OCR_START in text
    assert "DELVAUX" not in text and "Louise" not in text
